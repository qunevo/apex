"""Header-based extraction; scan saved rows beyond Excel table boundaries."""
from io import BytesIO
import re

from openpyxl import load_workbook

from .common import identifier, indexed, number, require

DISPATCH = {"operation id": "operation_id", "lot": "lot_id", "order": "order_id",
            "qty (pcs)": "quantity", "workplace": "machine_id", "sequence": "sequence",
            "person": "person_id", "planned start": "start", "setup (min)": "setup_minutes",
            "run (min)": "run_minutes", "fixed": "fixed", "attendance (min)": "attendance_minutes",
            "planner note": "note"}
SKILLS = ("Setup", "Mechanical", "Precision", "Electrical", "Calibration", "Testing")


def normalized(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def yes_no(value, location):
    require(normalized(value) in ("yes", "no"), "BOOLEAN", location, "Use Yes or No")
    return normalized(value) == "yes"


def table(book, sheet_name, columns, key_header):
    require(sheet_name in book.sheetnames, "SHEET", sheet_name, "Required worksheet is missing")
    sheet = book[sheet_name]
    # Do not trust stored table refs/dimensions: Excel users may append outside a table.
    sheet.reset_dimensions()
    headers, output = None, []
    for cells in sheet.iter_rows():
        labels = [normalized(cell.value) for cell in cells]
        if headers is None:
            if key_header not in labels:
                continue
            require(set(columns) <= set(labels), "HEADERS", sheet_name,
                    "Missing columns: " + ", ".join(sorted(set(columns) - set(labels))))
            require(all(labels.count(key) == 1 for key in columns), "HEADERS", sheet_name,
                    "Required column names must be unique")
            headers = {columns[label]: pos for pos, label in enumerate(labels) if label in columns}
            continue
        values = {key: cells[pos] if pos < len(cells) else None for key, pos in headers.items()}
        if all(cell is None or cell.value in (None, "") for cell in values.values()):
            continue
        row_number = next(cell.row for cell in cells if cell.value is not None)
        location = f"{sheet_name}!{row_number}"
        row = {key: cell.value if cell else None for key, cell in values.items()}
        for key, cell in values.items():
            require(cell is None or cell.data_type not in ("f", "e"), "EXCEL_INPUT", location,
                    f"{key} must be a saved input value, not a formula/error; calculated output columns are ignored")
        row["_row"] = location
        output.append(row)
    require(headers is not None, "HEADERS", sheet_name, "Could not find the header row")
    return output


def read_workbook(content):
    book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        dispatch = table(book, "Dispatch plan", DISPATCH, "operation id")
        skills = table(book, "Skills", {"person": "person_id", **{s.casefold(): s for s in SKILLS}}, "person")
        for row in dispatch:
            loc = row["_row"]
            for key in ("operation_id", "lot_id", "order_id"):
                row[key] = identifier(row[key], loc + ":" + key)
            row["fixed"] = yes_no(row["fixed"], loc + ":fixed")
            for key in ("quantity", "sequence", "setup_minutes", "run_minutes", "attendance_minutes"):
                if row[key] in (None, ""):
                    row[key] = None
                else:
                    row[key] = number(row[key], loc + ":" + key, positive=key in ("quantity", "sequence", "run_minutes"),
                                      integer=key in ("quantity", "sequence"))
            for key in ("machine_id", "person_id"):
                row[key] = identifier(row[key], loc + ":" + key) if row[key] not in (None, "") else None
            if row["start"] == "":
                row["start"] = None
            require(row["machine_id"] or all(row[k] is None for k in ("setup_minutes", "run_minutes", "attendance_minutes")),
                    "ALLOWANCE_MACHINE", loc, "Duration allowances need the machine they apply to")
            require(not row["fixed"] or all(row[k] not in (None, "") for k in ("machine_id", "person_id", "start")),
                    "FIXED_INPUT", loc, "Fixed rows require workplace, person and planned start")
        for row in skills:
            row["person_id"] = identifier(row["person_id"], row["_row"])
            for skill in SKILLS:
                row[skill] = yes_no(row[skill], row["_row"] + ":" + skill)
        return indexed(dispatch, "Dispatch plan", "operation_id"), indexed(skills, "Skills", "person_id")
    finally:
        book.close()
